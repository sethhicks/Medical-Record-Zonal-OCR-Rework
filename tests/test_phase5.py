# tests/test_phase5.py
"""Phase 5 tests — OUT-01 through OUT-05 (Excel export writer).

Writer produces two sheets (CMS-1500, UB-04) each with 3 columns:
  Patient Name | Total Charge | Date of Service
"""
import pytest
from pathlib import Path

_ROOT = Path(__file__).parent.parent


def _make_results(field_names, value="", confidence=95.0):
    from models.field_result import FieldResult
    return [FieldResult(field_name=name, value=value, confidence=confidence)
            for name in field_names]


# ---------------------------------------------------------------------------
# OUT-01 + OUT-02 + OUT-03: write_workbook — writer structure and file output
# ---------------------------------------------------------------------------

def test_write_workbook_returns_str(tmp_path):
    """write_workbook returns a str (file path to created workbook) (OUT-01)."""
    from pipeline import write_workbook
    result = write_workbook(
        cms_pages=[[]],
        ub_pages=[[]],
        settings={"output_dir": str(tmp_path)},
    )
    assert isinstance(result, str)


def test_output_file_created(tmp_path):
    """write_workbook creates the output .xlsx file on disk (OUT-01)."""
    from pipeline import write_workbook
    result = write_workbook(
        cms_pages=[[]],
        ub_pages=[[]],
        settings={"output_dir": str(tmp_path)},
    )
    assert Path(result).exists()


def test_two_sheets_named_correctly(tmp_path):
    """Workbook contains exactly two sheets: 'CMS-1500' and 'UB-04' (OUT-02)."""
    import openpyxl
    from pipeline import write_workbook
    result = write_workbook(
        cms_pages=[[]],
        ub_pages=[[]],
        settings={"output_dir": str(tmp_path)},
    )
    wb = openpyxl.load_workbook(result)
    assert wb.sheetnames == ["CMS-1500", "UB-04"]


def test_cms1500_column_count(tmp_path):
    """CMS-1500 sheet has exactly 3 columns (OUT-02)."""
    import openpyxl
    from pipeline import write_workbook
    results = _make_results(["patient_name", "total_charge", "date_of_service_sl1"])
    result = write_workbook(
        cms_pages=[results],
        ub_pages=[[]],
        settings={"output_dir": str(tmp_path)},
    )
    wb = openpyxl.load_workbook(result)
    assert wb["CMS-1500"].max_column == 3


def test_ub04_column_count(tmp_path):
    """UB-04 sheet has exactly 3 columns (OUT-02)."""
    import openpyxl
    from pipeline import write_workbook
    results = _make_results(["patient_name", "total_charge", "date_of_service_rl1"])
    result = write_workbook(
        cms_pages=[[]],
        ub_pages=[results],
        settings={"output_dir": str(tmp_path)},
    )
    wb = openpyxl.load_workbook(result)
    assert wb["UB-04"].max_column == 3


def test_cms1500_header_row_frozen(tmp_path):
    """CMS-1500 sheet header row is frozen (freeze_panes == 'A2') (OUT-03)."""
    import openpyxl
    from pipeline import write_workbook
    result = write_workbook(
        cms_pages=[[]],
        ub_pages=[[]],
        settings={"output_dir": str(tmp_path)},
    )
    wb = openpyxl.load_workbook(result)
    assert wb["CMS-1500"].freeze_panes == "A2"


def test_no_yellow_fill(tmp_path):
    """Cells are never highlighted — yellow fill has been removed (OUT-04)."""
    import openpyxl
    from pipeline import write_workbook
    results = _make_results(["total_charge"], value="150.00", confidence=10.0)
    result = write_workbook(
        cms_pages=[results],
        ub_pages=[[]],
        settings={"output_dir": str(tmp_path)},
    )
    wb = openpyxl.load_workbook(result)
    ws = wb["CMS-1500"]
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            fill = cell.fill
            assert fill is None or fill.fill_type == "none" or not fill.fgColor.rgb.endswith("FFFF00")


# ---------------------------------------------------------------------------
# OUT-01 + OUT-05: pipeline import and text formatting
# ---------------------------------------------------------------------------

def test_import_write_workbook_from_pipeline():
    """write_workbook is importable from the pipeline package (OUT-01)."""
    from pipeline import write_workbook
    assert callable(write_workbook)


def test_text_format_charge_column(tmp_path):
    """Total Charge column cells use text format '@' to prevent Excel number conversion (OUT-05)."""
    import openpyxl
    from pipeline import write_workbook
    results = _make_results(["total_charge"], value="123.45", confidence=95.0)
    result = write_workbook(
        cms_pages=[results],
        ub_pages=[[]],
        settings={"output_dir": str(tmp_path)},
    )
    wb = openpyxl.load_workbook(result)
    ws = wb["CMS-1500"]
    header_row = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    col_idx = header_row.index("Total Charge") + 1
    assert ws.cell(2, col_idx).number_format == "@"


def test_text_format_date_column(tmp_path):
    """Date of Service column cells use text format '@' to preserve date strings as-is (OUT-05)."""
    import openpyxl
    from pipeline import write_workbook
    results = _make_results(["date_of_service_sl1"], value="01/01/2024", confidence=95.0)
    result = write_workbook(
        cms_pages=[results],
        ub_pages=[[]],
        settings={"output_dir": str(tmp_path)},
    )
    wb = openpyxl.load_workbook(result)
    ws = wb["CMS-1500"]
    header_row = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    col_idx = header_row.index("Date of Service") + 1
    assert ws.cell(2, col_idx).number_format == "@"


def test_cms1500_one_row_per_page(tmp_path):
    """CMS-1500 sheet has one row per page (header + N data rows for N pages) (OUT-02)."""
    import openpyxl
    from pipeline import write_workbook
    page1 = _make_results(["patient_name"], value="Smith, John", confidence=90.0)
    page2 = _make_results(["patient_name"], value="Doe, Jane", confidence=90.0)
    result = write_workbook(
        cms_pages=[page1, page2],
        ub_pages=[[]],
        settings={"output_dir": str(tmp_path)},
    )
    wb = openpyxl.load_workbook(result)
    assert wb["CMS-1500"].max_row == 3  # header row + 2 data rows
